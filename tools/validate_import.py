"""
Standalone validator for the SRMAP_KRA-KPI workbook import.

Re-implements (in Python, against the raw .xlsx) the same header
normalization / header-row discovery / row-filtering logic that
src/03_ImportPipeline.gs uses in Apps Script, so the import can be sanity
checked -- KPI row counts per office/scorecard against 01_Index, Wt(%) sums,
Goal ID extraction -- WITHOUT needing to push to a live Google Sheet first.

Usage:
    pip install openpyxl
    python3 tools/validate_import.py /path/to/SRMAP_KRA-KPI_AY26-27.xlsx

If you change the header-matching logic in 03_ImportPipeline.gs (HEADER_ALIASES,
TIER01_KPI_FIELDS, TIER2_KPI_FIELDS, the "TOTAL row" filter, etc.), mirror the
change here and re-run against the real workbook before trusting a live import.
"""
import re
import sys
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else None
if not SRC:
    print("Usage: python3 validate_import.py /path/to/workbook.xlsx")
    sys.exit(1)
wb = openpyxl.load_workbook(SRC, data_only=True)

def normalize(s):
    return re.sub(r'[^a-z0-9]+', '', str(s if s is not None else '').lower())

HEADER_ALIASES = {
    'S.No': ['S.No', '#'],
    'Strategic Pillar': ['Strategic Pillar', 'Pillar'],
    'How to Measure (KPI-Specific Formula)': ['How to Measure (KPI-Specific Formula)', 'How it is measured'],
    'Q1 Milestone Target': ['Q1 Milestone Target', 'Q1'],
    'Q2 Milestone Target': ['Q2 Milestone Target', 'Q2'],
    'Q3 Milestone Target': ['Q3 Milestone Target', 'Q3'],
    'Q4 Milestone Target': ['Q4 Milestone Target', 'Q4'],
}

def build_lookup(canonical_fields):
    lookup = {}
    for c in canonical_fields:
        for v in HEADER_ALIASES.get(c, [c]):
            lookup[normalize(v)] = c
    return lookup

def find_header_row(ws, canonical_fields, max_scan=6):
    lookup = build_lookup(canonical_fields)
    best = (-1, 0, {})
    scan_rows = min(max_scan, ws.max_row)
    for r in range(1, scan_rows + 1):
        col_map = {}
        matches = 0
        for c in range(1, ws.max_column + 1):
            norm = normalize(ws.cell(row=r, column=c).value)
            if norm and norm in lookup:
                col_map[lookup[norm]] = c
                matches += 1
        if matches > best[1]:
            best = (r, matches, col_map)
    if best[1] < 2:
        return None
    return best

def read_rows(ws, header_row, col_map):
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        vals = {name: ws.cell(row=r, column=c).value for name, c in col_map.items()}
        if any(str(v).strip() for v in vals.values() if v is not None):
            rows.append(vals)
    return rows

def filter_kpi_rows(rows):
    out = []
    for r in rows:
        kpi = str(r.get('KPI') or '').strip()
        if not kpi or kpi.upper() == 'TOTAL':
            continue
        out.append(r)
    return out

def extract_goal_id(text):
    m = re.search(r'P\d+\.\d+', str(text or ''))
    return m.group(0) if m else ''

def parse_tier(text):
    m = re.search(r'\d+', str(text or ''))
    return m.group(0) if m else str(text or '').strip()

# ---- 01_Index ----
idx_ws = wb['01_Index']
idx_expected = ['Sheet', 'Full name', 'Tier', 'Cluster', 'KPIs', 'Strategic', 'Operational', 'Verticals', 'Total Wt']
idx_header = find_header_row(idx_ws, idx_expected)
print("01_Index header row:", idx_header[0], "matches:", idx_header[1])
idx_rows = read_rows(idx_ws, idx_header[0], idx_header[2])
idx_rows = [r for r in idx_rows if str(r.get('Sheet') or '').strip()]
print("01_Index data rows:", len(idx_rows))

index_by_sheet = {str(r['Sheet']).strip(): r for r in idx_rows}

# ---- 03_Strategic_Pillars ----
pil_ws = wb['03_Strategic_Pillars']
pil_expected = ['Pillar', 'Pillar name', 'Goal ID', 'Five-year goal / target', 'Owner office', 'Contributing offices']
pil_header = find_header_row(pil_ws, pil_expected)
print("\n03_Strategic_Pillars header row:", pil_header[0], "matches:", pil_header[1])
pil_rows = read_rows(pil_ws, pil_header[0], pil_header[2])
goal_ids = [str(r['Goal ID']).strip() for r in pil_rows if r.get('Goal ID')]
print("Goals found:", len(goal_ids), "unique:", len(set(goal_ids)))
pillars_seen = set()
cur_pillar = ''
for r in pil_rows:
    if r.get('Pillar'):
        cur_pillar = str(r['Pillar']).strip()
    pillars_seen.add(cur_pillar)
print("Pillars found via forward-fill:", sorted(pillars_seen))

# ---- 04_Cascade_Map ----
cm_ws = wb['04_Cascade_Map']
cm_expected = ['Cluster', 'Office', 'Goals owned', 'Goals contributed to', 'KPIs', 'Strategic', 'Operational']
cm_header = find_header_row(cm_ws, cm_expected)
print("\n04_Cascade_Map header row:", cm_header[0], "matches:", cm_header[1])
cm_rows = read_rows(cm_ws, cm_header[0], cm_header[2])
print("Cascade map office rows:", len(cm_rows))

# ---- Tier 0/1 scorecards ----
TIER0_1 = ['VC Scorecard', 'C1 Academic Scorecard', 'C2 Growth Scorecard', 'C3 Student Scorecard', 'C4 Governance Scorecard']
TIER01_FIELDS = ['S.No', 'Tier', 'Cluster', 'KRA Type', 'Strategic Pillar', '5-Yr Goal (ID + target)', 'Goal Role',
                  'Vertical', 'KRA', 'KPI', 'How to Measure (KPI-Specific Formula)', 'Unit', 'Frequency',
                  'Annual Target (AY 26-27)', 'Wt (%)', 'Rolls Up From',
                  'Q1 Milestone Target', 'Q2 Milestone Target', 'Q3 Milestone Target', 'Q4 Milestone Target', 'Owner Sign-off']

print("\n--- Tier 0/1 scorecards ---")
tier01_total = 0
for name in TIER0_1:
    ws = wb[name]
    header = find_header_row(ws, TIER01_FIELDS)
    if not header:
        print(name, "HEADER ROW NOT FOUND")
        continue
    rows = filter_kpi_rows(read_rows(ws, header[0], header[2]))
    wt_sum = sum(float(r.get('Wt (%)') or 0) for r in rows)
    tier01_total += len(rows)
    idx_count = index_by_sheet.get(name, {}).get('KPIs')
    print(f"{name}: header_row={header[0]} matches={header[1]} rows={len(rows)} (01_Index says {idx_count}) wt_sum={wt_sum:.2f}")

# ---- Tier 2 office scorecards ----
TIER2_FIELDS = ['S.No', 'Office', 'Tier', 'Cluster', 'KRA Type', 'Strategic Pillar', '5-Yr Goal (ID + target)', 'Goal Role',
                 'Vertical', 'KRA', 'KPI', 'Strategic Rationale (5-Yr Goals)', 'Definition of KPI',
                 'How to Measure (KPI-Specific Formula)', 'Illustration', 'Unit', 'Frequency', 'Measurement Nature',
                 'Metric Type', 'Indicator Nature', 'Wt (%)', 'Annual Target (AY 26-27)',
                 'Q1 Milestone Target', 'Q2 Milestone Target', 'Q3 Milestone Target', 'Q4 Milestone Target',
                 'Q1 Actual', 'Q2 Actual', 'Q3 Actual', 'Q4 Actual', 'Q1 Status', 'Q2 Status', 'Q3 Status', 'Q4 Status',
                 'Q1 ATR', 'Q2 ATR', 'Q3 ATR', 'Q4 ATR', 'Score (Actual vs Target)', 'Data Source', 'KPI Owner',
                 'R (Responsible)', 'A (Accountable)', 'C (Consulted)', 'I (Informed)', 'Responsible Officer',
                 'Escalation / Remarks', 'Overlap / Remarks']

print("\n--- Tier 2 office scorecards ---")
mismatches = []
weight_fails = []
tier2_total = 0
coowner_count = 0
missing_goal_id = 0
for sheet_name, idxrow in index_by_sheet.items():
    tier = parse_tier(idxrow.get('Tier'))
    if tier != '2':
        continue
    if sheet_name not in wb.sheetnames:
        print(sheet_name, "SHEET NOT FOUND IN WORKBOOK")
        continue
    ws = wb[sheet_name]
    header = find_header_row(ws, TIER2_FIELDS)
    if not header:
        print(sheet_name, "HEADER ROW NOT FOUND")
        continue
    rows = filter_kpi_rows(read_rows(ws, header[0], header[2]))
    tier2_total += len(rows)
    expected_count = idxrow.get('KPIs')
    wt_sum = sum(float(r.get('Wt (%)') or 0) for r in rows)
    for r in rows:
        if str(r.get('Goal Role') or '').strip().lower() == 'co-owner':
            coowner_count += 1
        gid = extract_goal_id(r.get('5-Yr Goal (ID + target)'))
        if not gid:
            missing_goal_id += 1
    status = "OK" if len(rows) == expected_count else "MISMATCH"
    if status == "MISMATCH":
        mismatches.append((sheet_name, len(rows), expected_count))
    if abs(wt_sum - 100) > 0.5:
        weight_fails.append((sheet_name, wt_sum))
    print(f"{sheet_name}: header_row={header[0]} rows={len(rows)} expected={expected_count} [{status}] wt_sum={wt_sum:.2f}")

print("\n=== SUMMARY ===")
print("Tier 0/1 KPI rows total:", tier01_total)
print("Tier 2 KPI rows total:", tier2_total)
print("Grand total KPI rows:", tier01_total + tier2_total)
print("Co-owner rows (-> KPI_Office_Links):", coowner_count)
print("Rows where Goal ID extraction FAILED:", missing_goal_id)
print("Office KPI-count mismatches vs 01_Index:", mismatches)
print("Offices whose Wt(%) sum is off by >0.5 from 100:", weight_fails)
